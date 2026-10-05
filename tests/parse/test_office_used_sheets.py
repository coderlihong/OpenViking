"""Real Office payload probes and process lifecycle checks for KB-012.

These deterministic synthetic fixtures are contract evidence, not quality gold.
"""
import asyncio
import json
import os
import subprocess
import sys
import time
import zipfile
from pathlib import Path

import pytest

from openviking.parse.office_converter import OfficeConverter, OfficeConverterConfig, OfficeConversionError
from openviking.parse.office_layout import prepare_workbook


def make_fixtures(root):
    from docx import Document
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image
    from PIL import Image as PILImage
    from pptx import Presentation
    from pptx.util import Inches
    root.mkdir(parents=True, exist_ok=True)
    picture = root / "diagram.png"
    PILImage.new("RGB", (120, 80), "red").save(picture)
    doc = Document()
    doc.add_heading("Repair Alpha", 1)
    doc.add_paragraph("WORD_HEAD temperature 42 degrees")
    doc.add_picture(str(picture))
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Pressure MPa"
    table.cell(1, 0).text = "WORD_TABLE 12.5"
    doc.add_heading("Repair Omega", 1)
    doc.add_paragraph("WORD_TAIL pressure 12.5")
    doc.save(root / "word.docx")
    wb = Workbook()
    sheet = wb.active
    sheet.title = "Main"
    sheet.append(["Pressure MPa", "Reading"])
    sheet.append(["EXCEL_HEAD", 12.5])
    sheet.merge_cells("A3:B3")
    sheet["A3"] = "MERGED_CELL"
    sheet.add_image(Image(str(picture)), "C4")
    sheet["A90"] = "EXCEL_TAIL"
    sheet.print_area = "A1:B2"  # Must not hide the long tail.
    hidden = wb.create_sheet("HiddenData")
    hidden["A1"] = "HIDDEN_SHEET_VALUE"
    hidden.sheet_state = "hidden"
    wb.create_sheet("Empty")
    wb.save(root / "excel.xlsx")
    deck = Presentation()
    for token in ["SLIDE_ALPHA", "SLIDE_OMEGA"]:
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text_frame.text = token
        slide.shapes.add_picture(str(picture), Inches(1), Inches(2), width=Inches(2))
        table = slide.shapes.add_table(2, 2, Inches(4), Inches(2), Inches(3), Inches(2)).table
        table.cell(0, 0).text = "Torque Nm"
        table.cell(1, 0).text = "SLIDE_TABLE 23"
    deck.save(root / "slides.pptx")
    return [root / "word.docx", root / "excel.xlsx", root / "slides.pptx"]


def test_xlsx_used_sheet_policy_preserves_payload(tmp_path):
    source = make_fixtures(tmp_path)[1]
    output = tmp_path / "normalized.xlsx"
    manifest = prepare_workbook(source, output)
    assert [(s["name"], s["printed"]) for s in manifest] == [("Main", True), ("HiddenData", True), ("Empty", False)]
    assert manifest[0]["used_range"] == "A1:D90"
    with zipfile.ZipFile(source) as before, zipfile.ZipFile(output) as after:
        assert set(before.namelist()) == set(after.namelist())
        for name in before.namelist():
            if name not in {"xl/workbook.xml", "xl/styles.xml"} and not name.startswith("xl/worksheets/sheet"):
                assert before.read(name) == after.read(name)
    from openpyxl import load_workbook
    rendered = load_workbook(output)
    assert rendered["Main"]["A90"].value == "EXCEL_TAIL"
    assert str(rendered["Main"].print_area).endswith("$A$1:$D$90")
    assert rendered["HiddenData"].sheet_state == "visible"


def test_resource_wrapper_enforces_cpu_and_file_limit():
    converter = OfficeConverter(OfficeConverterConfig(cpu_seconds=4, file_size_bytes=8192))
    probe = "import resource,json; print(json.dumps([resource.getrlimit(resource.RLIMIT_CPU),resource.getrlimit(resource.RLIMIT_FSIZE)]))"
    completed = subprocess.run(converter._limited_command([sys.executable, "-c", probe]), capture_output=True, text=True, check=True)
    assert json.loads(completed.stdout) == [[4, 4], [8192, 8192]]


@pytest.mark.asyncio
async def test_cancel_stops_real_child_and_cleans_owned_directory(tmp_path):
    from tests.parse.test_office_conversion import _minimal_docx
    source = _minimal_docx(tmp_path / "source.docx")
    pid_path = tmp_path / "pid"
    launcher = tmp_path / "soffice"
    launcher.write_text("#!/bin/sh\nif [ \"$1\" = \"--version\" ]; then echo 'LibreOffice synthetic'; exit 0; fi\necho $$ > " + str(pid_path) + "\nsleep 60\n")
    launcher.chmod(0o755)
    converter = OfficeConverter(OfficeConverterConfig(binary=str(launcher), temp_root=tmp_path / "work", ttl_seconds=0.01))
    task = asyncio.create_task(converter.convert(source))
    for _ in range(100):
        if pid_path.exists():
            break
        await asyncio.sleep(.02)
    assert pid_path.exists()
    time.sleep(.03)
    assert converter.cleanup_expired() == []  # Active conversion lease must survive TTL.
    pid = int(pid_path.read_text())
    task.cancel()
    with pytest.raises(asyncio.CancelledError):
        await task
    with pytest.raises(ProcessLookupError):
        os.kill(pid, 0)
    assert list(converter.temp_root.iterdir()) == []
    assert source.exists()


@pytest.mark.integration
@pytest.mark.skipif(not OfficeConverter().is_available(), reason="Actual LibreOffice is required")
def test_real_six_formats_preserve_text_tables_images_and_sheet_tail(tmp_path):
    import pymupdf
    sources = make_fixtures(tmp_path / "fixtures")
    converter = OfficeConverter(OfficeConverterConfig(temp_root=tmp_path / "work", check_fonts=False))
    binary = converter._locate_binary()
    for source, extension, filter_name in zip(sources, ["doc", "xls", "ppt"], ["MS Word 97", "MS Excel 97", "MS PowerPoint 97"]):
        completed = converter._run_libreoffice(source, source.parent, tmp_path / ("profile_" + extension), extension + ":" + filter_name, 60)
        assert completed.returncode == 0
        legacy = source.with_suffix("." + extension)
        assert legacy.is_file()
        sources.append(legacy)
        if len(sources) == 6:
            break
    for source in sources:
        result = converter.convert_sync(source)
        try:
            with pymupdf.open(result.pdf_path) as pdf:
                text = "\n".join(page.get_text() for page in pdf)
                expected = {"word": ["WORD_HEAD", "WORD_TABLE", "WORD_TAIL"],
                            "excel": ["EXCEL_HEAD", "EXCEL_TAIL", "MERGED_CELL", "HIDDEN_SHEET_VALUE"],
                            "presentation": ["SLIDE_ALPHA", "SLIDE_OMEGA", "SLIDE_TABLE"]}[result.kind]
                compact = "".join(text.split())
                assert all("".join(token.split()) in compact for token in expected), (source.suffix, text)
                assert any(page.get_images() for page in pdf), source.suffix
                if result.kind == "presentation":
                    assert len(pdf) == 2
                    assert "SLIDE_ALPHA" in pdf[0].get_text()
                    assert "SLIDE_OMEGA" in pdf[1].get_text()
                if result.kind == "excel":
                    assert len(pdf) > 2
                    assert result.structure["print_policy"] == "all_used_sheets_ignore_source_print_area"
            assert source.is_file()
        finally:
            converter.cleanup(result)
    assert list(converter.temp_root.iterdir()) == []
