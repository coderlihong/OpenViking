"""Live synthetic legacy Office contracts using the current server/SDK.

Conversion uses the frozen remote image and verified OOXML font declarations.
This is a transport/ingestion/evidence contract, not repair knowledge gold.
API credentials are read from the local test setup and never included in output.
"""
import argparse
import ast
import asyncio
import hashlib
import json
from pathlib import Path
import shlex
import subprocess
from types import SimpleNamespace
import sys
import zipfile
from xml.etree import ElementTree as ET

sys.path.insert(0,str(Path(__file__).resolve().parents[3]))
from tests.enterprise_kb.flows import flow_d1, flow_w1_pdf
from openviking.parse.office_converter import OfficeConverter
from openviking_sdk.client import AsyncHTTPClient

IMAGE='sha256:6a2cfcf932ad67e7f0445d50b87e446e0b5c9710053eee27324d54d06201cc70'
STATEMENT='Synthetic repair contract: SY215C hydraulic pump pressure is low. Check the inlet filter for blockage before replacing the hydraulic pump. The pressure reading is 12.5 MPa.'
REMOTE=r'''
import os, pathlib, shutil, subprocess, sys, tempfile
ext, filt=sys.argv[1:]
root=pathlib.Path(tempfile.mkdtemp(prefix='legacy_contract_',dir='/home/t-wanghai/models/openviking-office'))
try:
 src=root/('source'+ext);src.write_bytes(sys.stdin.buffer.read());src.chmod(0o400)
 out=root/'output';out.mkdir(mode=0o700)
 user=out/'profile'/'user';user.mkdir(parents=True)
 (user/'registrymodifications.xcu').write_text('<oor:items xmlns:oor="http://openoffice.org/2001/registry"><item oor:path="/org.openoffice.Office.Common/Security/Scripting"><prop oor:name="DisableMacrosExecution" oor:op="fuse"><value>true</value></prop><prop oor:name="DisableActiveContent" oor:op="fuse"><value>true</value></prop></item></oor:items>')
 cmd=['docker','run','--rm','--user',str(os.getuid())+':'+str(os.getgid()),'--network','none','--read-only','--cap-drop','ALL','--security-opt','no-new-privileges','--memory','4g','--cpus','2','--pids-limit','128','--tmpfs','/tmp:rw,nosuid,nodev,size=1g','-e','HOME=/tmp','-e','MAX_CONCURRENCY=2','-v',str(src)+':/input/source'+ext+':ro','-v',str(out)+':/output:rw','--entrypoint','/usr/bin/soffice','IMAGE','--headless','--norestore','--nofirststartwizard','-env:UserInstallation=file:///output/profile','--convert-to',filt,'--outdir','/output','/input/source'+ext]
 run=subprocess.run(cmd,capture_output=True,timeout=120)
 if run.returncode:sys.stderr.buffer.write(run.stderr);sys.exit(run.returncode)
 result=next(p for p in out.iterdir() if p.is_file())
 sys.stdout.buffer.write(result.read_bytes())
finally:shutil.rmtree(root)
'''.replace("'IMAGE'",repr(IMAGE))

def sha(path):return hashlib.sha256(path.read_bytes()).hexdigest()

def read_key(path):
 for node in ast.walk(ast.parse(path.read_text())):
  if isinstance(node,ast.Assign) and any(isinstance(target,ast.Name) and target.id=='KEY' for target in node.targets):
   if isinstance(node.value,ast.Constant) and isinstance(node.value.value,str):return node.value.value
 raise ValueError('Local KEY literal not found')

def pin_fonts(path):
 with zipfile.ZipFile(path) as source:parts={name:source.read(name) for name in source.namelist()}
 for name,blob in list(parts.items()):
  if not name.endswith('.xml') or not name.startswith(('word/','ppt/','xl/')):continue
  root=ET.fromstring(blob)
  for element in root.iter():
   tag=element.tag.rsplit('}',1)[-1]
   for key,value in list(element.attrib.items()):
    attr=key.rsplit('}',1)[-1]
    if (tag in {'rFonts','font','latin','ea','cs'} and attr in {'typeface','ascii','hAnsi','eastAsia','cs','name'}) or (name.startswith('xl/') and tag=='name' and attr=='val'):
     if value and not value.startswith(('+mn-','+mj-')):element.attrib[key]='Noto Sans CJK SC'
  parts[name]=ET.tostring(root,encoding='utf-8',xml_declaration=True)
 with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as destination:
  for name,blob in parts.items():destination.writestr(name,blob)
 assert OfficeConverter._declared_fonts(path,path.suffix)=={'Noto Sans CJK SC'}

def prepare(directory,tag):
 from docx import Document
 from pptx import Presentation
 from pptx.util import Inches
 from openpyxl import load_workbook
 base=Path('/tmp/ov-flow/office-I047')
 source_paths={'.docx':base/'w1-office-word-fontchecked-I047.docx','.xlsx':base/'w1-office-sheets-diagnostic-I047.xlsx','.pptx':base/'w1-office-slides-fontchecked-I047.pptx'}
 records=[]
 for ext,original in source_paths.items():
  source=directory/('source'+ext)
  if ext=='.docx':
   doc=Document(original);doc.add_paragraph(STATEMENT);doc.save(source)
  elif ext=='.xlsx':
   book=load_workbook(original);sheet=book.create_sheet('LegacyContract');sheet['A1']=STATEMENT;sheet.column_dimensions['A'].width=100;book.save(source)
  else:
   slides=Presentation(original);slides.slides[0].shapes.add_textbox(Inches(.5),Inches(5),Inches(9),Inches(1.5)).text_frame.text=STATEMENT;slides.save(source)
  pin_fonts(source)
  legacy,export={'.docx':('.doc','doc:MS Word 97'),'.xlsx':('.xls','xls:MS Excel 97'),'.pptx':('.ppt','ppt:MS PowerPoint 97')}[ext]
  command=' '.join(shlex.quote(value) for value in ['python3','-c',REMOTE,ext,export])
  converted=subprocess.run(['ssh','-o','BatchMode=yes','-o','ConnectTimeout=8','t-wanghai@10.147.254.3',command],input=source.read_bytes(),capture_output=True,timeout=150)
  if converted.returncode:raise RuntimeError('Legacy fixture conversion failed: '+converted.stderr.decode(errors='replace'))
  assert converted.stdout.startswith(b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1')
  for flow in ['d1','w1']:
   output=directory/(f'legacy-{flow}-{tag}-{legacy[1:]}'+legacy);output.write_bytes(converted.stdout)
   records.append({'family':legacy,'flow':flow,'original_path':str(original),'original_sha256':sha(original),'ooxml_path':str(source),'ooxml_sha256':sha(source),'declared_ooxml_fonts':['Noto Sans CJK SC'],'path':str(output),'legacy_sha256':sha(output),'synthetic_statement':STATEMENT,'legacy_font_limitation':'font inventory validated before LibreOffice export; legacy CFB source font introspection is not implemented'})
 return records

async def run(args):
 key=read_key(Path(args.key_setup));directory=Path(args.directory);directory.mkdir(parents=True,exist_ok=True)
 fixtures_path=directory/'fixtures.json'
 fixtures=json.loads(fixtures_path.read_text()) if fixtures_path.exists() else prepare(directory,args.tag)
 fixtures_path.write_text(json.dumps(fixtures,ensure_ascii=False,indent=2))
 report={'scope':'Legacy DOC/XLS/PPT real HTTP and Async SDK D1/W1 contracts','sample_class':'synthetic_not_gold','quality_gold':'NOT_RUN','tag':args.tag,'image':IMAGE,'fixtures':fixtures,'runs':[]}
 for fixture in fixtures:
  path=Path(fixture['path']);flow=fixture['flow'];label=flow+fixture['family'];result_path=directory/(label.replace('.','-')+'.json')
  if result_path.exists():
   saved=json.loads(result_path.read_text())
   if saved.get('passed') is True:
    report['runs'].append(saved);continue
  try:
   if flow=='d1':
    result=await flow_d1.run(SimpleNamespace(pdf=str(path),url=args.url,key=key,account=None,user=None,backend='mineru',timeout=1800,query='SY215C hydraulic pump low pressure inlet filter blockage 12.5 MPa'))
   else:
    result=await flow_w1_pdf.run(SimpleNamespace(pdf=str(path),url=args.url,key=key,collection='sany',job_id=None,build_id=None,ingest_key=f'{args.tag}-legacy-{fixture["family"]}-ingest',exec_key=f'{args.tag}-legacy-{fixture["family"]}-exec',timeout=1800,deadline_seconds=1800,poll_attempts=None,poll_interval=3,target_prefix='viking://resources/wiki/sany',out=None))
    if result.get('assert_draft_ready'):
     client=AsyncHTTPClient(url=args.url,api_key=key,timeout=1800);await client.initialize()
     try:
      final=await client.kb.get_ingestion_job(result['job_id']);receipts=(final.get('checkpoint') or {}).get('draft_receipt',{}).get('receipts',[])
      result['candidate_evidence']=[]
      for index in range(len(receipts)):
       draft=await client.kb.get_wiki_draft(result['job_id'],index)
       for candidate in draft.get('candidates',[]):result['candidate_evidence'].append({'uri':candidate.get('uri'),'evidence':candidate.get('evidence'),'contributors':candidate.get('contributors')})
     finally:await client.close()
   result.update({'family':fixture['family'],'flow':flow,'legacy_sha256':fixture['legacy_sha256']})
  except Exception as error:result={'family':fixture['family'],'flow':flow,'passed':False,'error_type':type(error).__name__,'error':str(error)}
  result_path.write_text(json.dumps(result,ensure_ascii=False,indent=2,default=str));report['runs'].append(result)
  Path(args.out).write_text(json.dumps(report,ensure_ascii=False,indent=2,default=str))
  print(label+': '+str(result.get('passed')),flush=True)
 report['contract_passed']=len(report['runs'])==6 and all(result.get('passed') is True for result in report['runs'])
 report['legacy_binary_font_validation']='NOT_RUN; source OOXML declaration inventory was normalized before frozen LibreOffice export'
 report['office_feature_acceptance']='NOT_RUN'
 report['assert_visible_xls_image_materialized']=None
 image_probe=directory/'xls-image-probe.json'
 if image_probe.exists():
  probe=json.loads(image_probe.read_text())
  expected=probe['pdf'].get('displayed_image_count',0)
  xls=next((result for result in report['runs'] if result.get('flow')=='d1' and result.get('family')=='.xls'),{})
  if expected:
   report['assert_visible_xls_image_materialized']='images' in xls.get('tree_names',[])
   report['known_image_defect']={'visible_pdf_image_count':expected,'stored_xls_tree':xls.get('tree_names',[]),'source_probe':str(image_probe),'status':'FAIL' if not report['assert_visible_xls_image_materialized'] else 'PASS'}
 report['passed']=report['contract_passed'] and report['assert_visible_xls_image_materialized'] is not False
 correction=directory/'harness-path-correction.json'
 if correction.exists():report['harness_correction']=json.loads(correction.read_text())
 Path(args.out).write_text(json.dumps(report,ensure_ascii=False,indent=2,default=str))
 return report

if __name__=='__main__':
 parser=argparse.ArgumentParser();parser.add_argument('--url',default='http://127.0.0.1:1933');parser.add_argument('--tag',default='I050');parser.add_argument('--key-setup',default='/tmp/ov-flow/probe_w1.py');parser.add_argument('--directory',default='/tmp/ov-flow/office-legacy-I050');parser.add_argument('--out',default='docs/spec/evidence/flow-office-legacy-I050-2026-10-05.json')
 report=asyncio.run(run(parser.parse_args()));sys.exit(0 if report['passed'] else 1)
